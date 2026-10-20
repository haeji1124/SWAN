"""관리자 화면 API (Phase 2): 기수, 신청자, 졸업생 · (Phase 6) 관리자 추가·해제. 모든 요청은 관리자만 (auth.require_admin).

- 기수: 목록 · 추가 · 수정 · 모집 열기/닫기(모집 중은 하나만) · 수료 처리(신청자를 졸업생 명단으로)
- 신청자: 기수별 표 · 엑셀로 내려받기 (취소한 사람은 따로 표시)
- 졸업생: 기수별 표 · 엑셀로 올리기(미리보기 → 확인 → 올리기)

엑셀 파일은 요청 본문 그대로 받는다 (multipart 패키지 없이, 사내망 반입 패키지를 늘리지 않으려고).
"""
import io
from datetime import date, datetime
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Request, Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

import auth
import quiz_api
from apply_api import PYTHON_EXP, cohort_info
from db import get_db
from models import Alumni, Application, Cohort, User

router = APIRouter(prefix="/api/admin", dependencies=[Depends(auth.require_admin)])

STATUS_LABEL = {"before": "모집 전", "open": "모집 중", "closed": "모집 마감", "done": "종료"}
MAX_UPLOAD = 2 * 1024 * 1024  # 졸업생 엑셀은 2MB까지


# ---------- 기수 ----------

def counts(db: Session, cohort_id: int) -> tuple:
    rows = db.execute(
        select(Application.canceled_at.is_(None), func.count()).where(Application.cohort_id == cohort_id).group_by(Application.canceled_at.is_(None))
    ).all()
    active = sum(n for ok, n in rows if ok)
    canceled = sum(n for ok, n in rows if not ok)
    return active, canceled


@router.get("/cohorts")
def cohorts(db: Session = Depends(get_db)):
    out = []
    for c in db.scalars(select(Cohort).order_by(Cohort.number.desc())).all():
        active, canceled = counts(db, c.id)
        alumni = db.scalar(select(func.count()).select_from(Alumni).where(Alumni.cohort_number == c.number))
        out.append({**cohort_info(c), "status_label": STATUS_LABEL.get(c.status, c.status),
                    "applicants": active, "canceled": canceled, "alumni": alumni})
    return out


class CohortIn(BaseModel):
    number: int = Field(ge=1, le=999)
    recruit_start: date
    recruit_end: date
    start_date: date
    end_date: date
    class_day: str = Field(min_length=1, max_length=20)
    class_time: str = Field(min_length=1, max_length=20)


def check_dates(body: CohortIn) -> None:
    # 화면에 그대로 보여 줄 한 문장으로 돌려준다 (Pydantic 검증 오류는 목록이라 화면에서 읽기 어렵다)
    if not (body.recruit_start <= body.recruit_end <= body.start_date <= body.end_date):
        raise HTTPException(status_code=400, detail="날짜 순서를 확인해 주세요: 모집 시작 ≤ 모집 마감 ≤ 개강 ≤ 종료")


def number_taken(db: Session, number: int, except_id: Optional[int] = None) -> bool:
    q = select(Cohort.id).where(Cohort.number == number)
    if except_id:
        q = q.where(Cohort.id != except_id)
    return db.scalars(q).first() is not None


@router.post("/cohorts")
def add_cohort(body: CohortIn, db: Session = Depends(get_db)):
    check_dates(body)
    if number_taken(db, body.number):
        raise HTTPException(status_code=409, detail=f"{body.number}기는 이미 있어요")
    c = Cohort(**body.model_dump(), status="before")
    db.add(c)
    db.commit()
    return cohort_info(c)


def get_cohort(db: Session, cohort_id: int) -> Cohort:
    c = db.get(Cohort, cohort_id)
    if c is None:
        raise HTTPException(status_code=404, detail="없는 기수예요")
    return c


@router.put("/cohorts/{cohort_id}")
def edit_cohort(cohort_id: int, body: CohortIn, db: Session = Depends(get_db)):
    c = get_cohort(db, cohort_id)
    check_dates(body)
    if number_taken(db, body.number, except_id=c.id):
        raise HTTPException(status_code=409, detail=f"{body.number}기는 이미 있어요")
    for k, v in body.model_dump().items():
        setattr(c, k, v)
    db.commit()
    return cohort_info(c)


@router.post("/cohorts/{cohort_id}/open")
def open_recruit(cohort_id: int, db: Session = Depends(get_db)):
    c = get_cohort(db, cohort_id)
    other = auth.open_cohort(db)
    if other and other.id != c.id:
        raise HTTPException(status_code=409, detail=f"{other.number}기 모집이 열려 있어요. 먼저 닫아 주세요 (모집 중인 기수는 하나만)")
    if c.status == "done":
        raise HTTPException(status_code=409, detail="종료된 기수는 다시 열 수 없어요")
    c.status = "open"
    db.commit()
    return cohort_info(c)


@router.post("/cohorts/{cohort_id}/close")
def close_recruit(cohort_id: int, db: Session = Depends(get_db)):
    c = get_cohort(db, cohort_id)
    if c.status != "open":
        raise HTTPException(status_code=409, detail="모집 중인 기수만 닫을 수 있어요")
    c.status = "closed"
    db.commit()
    return cohort_info(c)


@router.post("/cohorts/{cohort_id}/graduate")
def graduate(cohort_id: int, db: Session = Depends(get_db)):
    """수료 처리: 이 기수의 유효한 신청자를 졸업생 명단으로 옮기고 기수를 종료한다. 이미 명단에 있는 사람은 건너뛴다."""
    c = get_cohort(db, cohort_id)
    if c.status == "open":
        raise HTTPException(status_code=409, detail="모집 중인 기수는 수료 처리할 수 없어요. 먼저 모집을 닫아 주세요")
    apps = db.execute(
        select(Application, User).join(User, User.id == Application.user_id)
        .where(Application.cohort_id == c.id, Application.canceled_at.is_(None))
    ).all()
    moved = skipped = 0
    for a, u in apps:
        a.completed = True
        exists = db.scalars(select(Alumni.id).where(Alumni.emp_no == u.emp_no, Alumni.cohort_number == c.number)).first()
        if exists:
            skipped += 1
            continue
        db.add(Alumni(name=u.name, dept=u.dept, cohort_number=c.number, emp_no=u.emp_no))
        moved += 1
    c.status = "done"
    db.commit()
    return {"moved": moved, "skipped": skipped}


# ---------- 신청자 ----------

def applicant_rows(db: Session, c: Cohort) -> list:
    rows = db.execute(
        select(Application, User).join(User, User.id == Application.user_id)
        .where(Application.cohort_id == c.id).order_by(Application.applied_at)
    ).all()
    scores = quiz_api.first_scores(db)
    return [
        {
            "id": a.id, "name": u.name, "dept": u.dept, "emp_no": u.emp_no,
            "applied_at": a.applied_at.isoformat(timespec="minutes"),
            "python_exp": a.python_exp, "python_exp_label": PYTHON_EXP.get(a.python_exp, a.python_exp),
            "want_to_make": a.want_to_make,
            "canceled_at": a.canceled_at.isoformat(timespec="minutes") if a.canceled_at else None,
            "quiz_first_score": scores.get(u.id),  # 수준 테스트 첫 응시 점수 (안 풀었으면 None)
        }
        for a, u in rows
    ]


@router.get("/applicants")
def applicants(cohort_id: int, db: Session = Depends(get_db)):
    c = get_cohort(db, cohort_id)
    rows = applicant_rows(db, c)
    return {
        "cohort": cohort_info(c),
        "active": [r for r in rows if not r["canceled_at"]],
        "canceled": [r for r in rows if r["canceled_at"]],
        "quiz_total": len(quiz_api.questions(db)),  # '수준 테스트' 칸을 "5 / 6"으로 보여 주려고
    }


@router.get("/applicants.xlsx")
def applicants_xlsx(cohort_id: int, db: Session = Depends(get_db)):
    c = get_cohort(db, cohort_id)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{c.number}기 신청자"
    ws.append(["이름", "부서", "사번", "신청 일시", "파이썬 경험", "만들고 싶은 것", "수준 테스트(첫 응시)", "상태", "취소 일시"])
    for r in applicant_rows(db, c):
        ws.append([
            r["name"], r["dept"], r["emp_no"], r["applied_at"].replace("T", " "), r["python_exp_label"],
            r["want_to_make"] or "", r["quiz_first_score"] if r["quiz_first_score"] is not None else "",
            "취소" if r["canceled_at"] else "신청", (r["canceled_at"] or "").replace("T", " "),
        ])
    for col, width in zip("ABCDEFGHI", (10, 14, 10, 18, 18, 40, 16, 8, 18)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    name = f"SWAN_{c.number}기_신청자_{date.today():%Y%m%d}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"},
    )


# ---------- 졸업생 ----------

@router.get("/alumni")
def alumni(cohort_number: Optional[int] = None, db: Session = Depends(get_db)):
    by_cohort = dict(db.execute(select(Alumni.cohort_number, func.count()).group_by(Alumni.cohort_number)).all())
    q = select(Alumni).order_by(Alumni.cohort_number.desc(), Alumni.id)
    if cohort_number:
        q = q.where(Alumni.cohort_number == cohort_number)
    return {
        "counts": [{"cohort_number": k, "count": v} for k, v in sorted(by_cohort.items())],
        "total": sum(by_cohort.values()),
        "rows": [{"id": a.id, "name": a.name, "dept": a.dept, "cohort_number": a.cohort_number, "emp_no": a.emp_no}
                 for a in db.scalars(q).all()],
    }


HEADERS = {"이름": "name", "부서": "dept", "기수": "cohort_number", "사번": "emp_no"}


async def read_alumni_xlsx(request: Request, db: Session) -> dict:
    """엑셀(첫 시트, 첫 줄은 열 이름: 이름 · 부서 · 기수 · 사번)을 읽어 올릴 줄과 문제 있는 줄로 나눈다."""
    raw = await request.body()
    if not raw:
        raise HTTPException(status_code=400, detail="파일이 비어 있어요")
    if len(raw) > MAX_UPLOAD:
        raise HTTPException(status_code=413, detail="파일이 너무 커요 (2MB까지)")
    try:
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
    except Exception:
        raise HTTPException(status_code=400, detail="엑셀 파일(.xlsx)을 읽지 못했어요")
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    missing = [h for h in HEADERS if h not in header]
    if missing:
        raise HTTPException(status_code=400, detail=f"첫 줄에 열 이름이 없어요: {', '.join(missing)} (이름 · 부서 · 기수 · 사번)")
    idx = {HEADERS[h]: header.index(h) for h in HEADERS}
    existing = set(db.execute(select(Alumni.emp_no, Alumni.cohort_number)).all())
    ok, errors, duplicates, seen = [], [], [], set()
    for line, values in enumerate(rows, start=2):
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        get = lambda k: values[idx[k]] if idx[k] < len(values) else None  # noqa: E731
        name, dept, emp = (str(get(k) or "").strip() for k in ("name", "dept", "emp_no"))
        try:
            number = int(str(get("cohort_number")).replace("기", "").strip())
        except (TypeError, ValueError):
            number = None
        problem = [m for m, bad in (("이름", not name), ("부서", not dept), ("사번", not emp), ("기수", not number)) if bad]
        if problem:
            errors.append({"line": line, "message": f"{', '.join(problem)}이 비었거나 잘못됐어요"})
            continue
        row = {"name": name, "dept": dept, "cohort_number": number, "emp_no": emp}
        if (emp, number) in existing or (emp, number) in seen:
            duplicates.append({"line": line, **row})
            continue
        seen.add((emp, number))
        ok.append(row)
    return {"rows": ok, "errors": errors, "duplicates": duplicates}


@router.post("/alumni/preview")
async def alumni_preview(request: Request, db: Session = Depends(get_db)):
    """올리기 전 미리보기: 올릴 줄, 이미 있는 사람(건너뜀), 문제 있는 줄."""
    return await read_alumni_xlsx(request, db)


@router.post("/alumni/import")
async def alumni_import(request: Request, db: Session = Depends(get_db)):
    parsed = await read_alumni_xlsx(request, db)
    for r in parsed["rows"]:
        db.add(Alumni(**r))
    db.commit()
    return {"added": len(parsed["rows"]), "skipped": len(parsed["duplicates"]), "errors": len(parsed["errors"])}


# ---------- 관리자 (Phase 6) ----------
# 목록 · 추가(사번이나 이름으로 찾기) · 해제. 자기 자신은 해제할 수 없다 (관리자가 0명이 되는 것도 이것으로 막힌다)

def admin_row(db: Session, u: User) -> dict:
    return {"id": u.id, "emp_no": u.emp_no, "name": u.name, "dept": u.dept, "type": auth.user_type(db, u)}


@router.get("/admins")
def admins(me: User = Depends(auth.require_admin), db: Session = Depends(get_db)):
    rows = db.scalars(select(User).where(User.is_admin.is_(True)).order_by(User.name)).all()
    return {"items": [{**admin_row(db, u), "me": u.id == me.id} for u in rows]}


@router.get("/users/search")
def search_users(q: str = "", db: Session = Depends(get_db)):
    """관리자가 아닌 사용자를 사번이나 이름으로 찾는다 (사이트에 한 번이라도 들어온 사람만 User에 있다)"""
    q = q.strip()
    if len(q) < 1:
        return {"items": []}
    rows = db.scalars(
        select(User).where(User.is_admin.is_(False), (User.emp_no.contains(q)) | (User.name.contains(q))).order_by(User.name).limit(10)
    ).all()
    return {"items": [admin_row(db, u) for u in rows]}


def get_user(db: Session, uid: int) -> User:
    u = db.get(User, uid)
    if u is None:
        raise HTTPException(status_code=404, detail="없는 사용자예요")
    return u


@router.post("/admins/{uid}")
def add_admin(uid: int, db: Session = Depends(get_db)):
    u = get_user(db, uid)
    if u.is_admin:
        raise HTTPException(status_code=409, detail="이미 관리자예요")
    u.is_admin = True
    db.commit()
    return admin_row(db, u)


@router.delete("/admins/{uid}")
def remove_admin(uid: int, me: User = Depends(auth.require_admin), db: Session = Depends(get_db)):
    if uid == me.id:
        raise HTTPException(status_code=409, detail="자기 자신은 관리자에서 해제할 수 없어요")
    u = get_user(db, uid)
    if not u.is_admin:
        raise HTTPException(status_code=409, detail="관리자가 아니에요")
    u.is_admin = False
    db.commit()
    return {"ok": True}
