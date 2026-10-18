"""Phase 2 수용 기준을 서버에서 확인한다: 신청 · 취소 · 중복 막기 · 모집 닫기 · 관리자 권한 · 엑셀 내려받기/올리기 · 수료 처리.

임시 데이터베이스(SWAN_DB)에 가상 데이터를 새로 넣고 돌린다. 개발용 swan.db는 건드리지 않는다.
실행: .venv/Scripts/python.exe -m pytest tests -q
"""
import io
import os
import sys
import tempfile

DB_FILE = os.path.join(tempfile.mkdtemp(), "test.db")
os.environ["SWAN_DB"] = DB_FILE
os.environ["SWAN_DEV_LOGIN"] = "1"
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

import seed  # noqa: E402
from main import app  # noqa: E402

# seed.py의 가상 사용자 순서: 1~5 직원, 6·7 5기 신청자, 8~10 졸업생, 11·12 관리자
EMPLOYEE, EMPLOYEE2, APPLICANT, ADMIN = 1, 2, 6, 11


@pytest.fixture(autouse=True)
def fresh_db():
    seed.run()
    yield


def as_user(uid):
    c = TestClient(app)
    c.cookies.set("swan_dev_user", str(uid) if uid else "fail")
    return c


def five(c):
    return next(x for x in c.get("/api/admin/cohorts").json() if x["number"] == 5)


def test_apply_then_button_state_then_cancel_then_reapply():
    u = as_user(EMPLOYEE)
    assert u.get("/api/me").json()["applied_current"] is False
    r = u.post("/api/apply", json={"python_exp": "script", "want_to_make": "  설비 로그 정리  "})
    assert r.status_code == 200
    done = r.json()
    assert done["cohort_number"] == 5 and done["name"] == "한서준" and done["want_to_make"] == "설비 로그 정리"
    assert u.get("/api/me").json()["applied_current"] is True  # 버튼이 [신청 완료]로
    assert u.post("/api/apply/cancel").status_code == 200
    assert u.get("/api/me").json()["applied_current"] is False  # 다시 [신청하기]로
    assert u.post("/api/apply", json={"python_exp": "none"}).status_code == 200  # 취소 뒤 다시 신청 가능


def test_same_person_cannot_apply_twice():
    u = as_user(EMPLOYEE)
    assert u.post("/api/apply", json={"python_exp": "regular"}).status_code == 200
    r = u.post("/api/apply", json={"python_exp": "regular"})
    assert r.status_code == 409 and r.json()["detail"] == "이미 신청하셨어요"


def test_only_two_fields_and_validation():
    u = as_user(EMPLOYEE)
    assert u.post("/api/apply", json={"python_exp": "expert"}).status_code == 422
    assert u.post("/api/apply", json={"python_exp": "none", "want_to_make": "가" * 301}).status_code == 422


def test_logged_out_cannot_apply():
    r = as_user(None).post("/api/apply", json={"python_exp": "none"})
    assert r.status_code == 401 and r.json()["detail"] == "로그인이 필요해요"


def test_closing_recruit_switches_everyone_to_interest():
    admin = as_user(ADMIN)
    assert admin.post(f"/api/admin/cohorts/{five(admin)['id']}/close").status_code == 200
    u = as_user(EMPLOYEE)
    me = u.get("/api/me").json()
    assert me["recruiting"] is False  # 상단 버튼이 [다음 기수 알림 받기]로
    assert u.post("/api/apply", json={"python_exp": "none"}).status_code == 409
    assert u.post("/api/interest").json() == {"ok": True, "already": False}
    assert u.post("/api/interest").json()["already"] is True
    assert u.get("/api/me").json()["interested"] is True  # [알림 신청됨]
    assert u.get("/api/home").json()["recruit"] is None


def test_only_one_open_cohort():
    admin = as_user(ADMIN)
    r = admin.post("/api/admin/cohorts", json={
        "number": 6, "recruit_start": "2027-03-01", "recruit_end": "2027-03-12",
        "start_date": "2027-04-01", "end_date": "2027-06-24", "class_day": "목", "class_time": "16:00"})
    assert r.status_code == 200
    six = r.json()["id"]
    r = admin.post(f"/api/admin/cohorts/{six}/open")
    assert r.status_code == 409 and "먼저 닫아" in r.json()["detail"]
    admin.post(f"/api/admin/cohorts/{five(admin)['id']}/close")
    assert admin.post(f"/api/admin/cohorts/{six}/open").status_code == 200


def test_cohort_date_order_and_duplicate_number():
    admin = as_user(ADMIN)
    bad = {"number": 6, "recruit_start": "2027-03-12", "recruit_end": "2027-03-01",
           "start_date": "2027-04-01", "end_date": "2027-06-24", "class_day": "목", "class_time": "16:00"}
    r = admin.post("/api/admin/cohorts", json=bad)
    assert r.status_code == 400 and "날짜 순서" in r.json()["detail"]
    r = admin.post("/api/admin/cohorts", json={**bad, "number": 5, "recruit_start": "2027-03-01", "recruit_end": "2027-03-12"})
    assert r.status_code == 409


def test_non_admin_blocked():
    for uid in (EMPLOYEE, APPLICANT):
        assert as_user(uid).get("/api/admin/cohorts").status_code == 403
    assert as_user(None).get("/api/admin/cohorts").status_code == 401


def test_applicants_table_and_excel_download():
    as_user(EMPLOYEE).post("/api/apply", json={"python_exp": "none", "want_to_make": "보고서 자동화"})
    e2 = as_user(EMPLOYEE2)
    e2.post("/api/apply", json={"python_exp": "regular"})
    e2.post("/api/apply/cancel")
    admin = as_user(ADMIN)
    cid = five(admin)["id"]
    data = admin.get(f"/api/admin/applicants?cohort_id={cid}").json()
    assert len(data["active"]) == 3  # 가상 신청자 2 + 한서준
    assert [r["name"] for r in data["canceled"]] == ["오지유"]  # 취소한 사람은 따로
    r = admin.get(f"/api/admin/applicants.xlsx?cohort_id={cid}")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:5] == ("이름", "부서", "사번", "신청 일시", "파이썬 경험")
    assert len(rows) == 1 + 4 and any(r[7] == "취소" for r in rows[1:])


def xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_alumni_upload_preview_then_import_changes_home_numbers():
    admin = as_user(ADMIN)
    before = admin.get("/api/home").json()["stats"]
    content = xlsx([
        ["이름", "부서", "기수", "사번"],
        ["가상일", "가상팀 A", 5, "Z00001"],
        ["가상이", "가상팀 B", "5기", "Z00002"],
        ["", "가상팀 C", 5, "Z00003"],  # 이름 빠짐 → 문제 있는 줄
        ["가상일", "가상팀 A", 5, "Z00001"],  # 같은 파일 안 중복 → 건너뜀
    ])
    p = admin.post("/api/admin/alumni/preview", content=content).json()
    assert len(p["rows"]) == 2 and len(p["errors"]) == 1 and len(p["duplicates"]) == 1
    r = admin.post("/api/admin/alumni/import", content=content).json()
    assert r["added"] == 2
    after = admin.get("/api/home").json()
    assert after["stats"]["alumni"] == before["alumni"] + 2
    assert after["stats"]["depts"] == before["depts"] + 2  # 부서 띠도 바뀐다
    assert "가상이" in after["alumni_names"]  # 이름 띠 (최신 기수가 앞)
    # 같은 파일을 또 올리면 모두 건너뛴다
    assert admin.post("/api/admin/alumni/import", content=content).json()["added"] == 0


def test_alumni_upload_rejects_wrong_file():
    admin = as_user(ADMIN)
    assert admin.post("/api/admin/alumni/preview", content=b"not excel").status_code == 400
    r = admin.post("/api/admin/alumni/preview", content=xlsx([["성명", "팀"], ["a", "b"]]))
    assert r.status_code == 400 and "열 이름" in r.json()["detail"]


def test_graduate_moves_applicants_to_alumni():
    admin = as_user(ADMIN)
    cid = five(admin)["id"]
    assert admin.post(f"/api/admin/cohorts/{cid}/graduate").status_code == 409  # 모집 중이면 안 됨
    admin.post(f"/api/admin/cohorts/{cid}/close")
    r = admin.post(f"/api/admin/cohorts/{cid}/graduate").json()
    assert r == {"moved": 2, "skipped": 0}
    assert five(admin)["status"] == "done" and five(admin)["alumni"] == 2
    assert as_user(APPLICANT).get("/api/me").json()["type"] == "alumni"  # 수료하면 졸업생


def test_home_shows_start_date():
    rec = as_user(EMPLOYEE).get("/api/home").json()["recruit"]
    assert rec["start_label"] == "11월 5일(목)" and rec["class_day"] == "목" and rec["class_time"] == "16:00"
